# -*- coding: utf-8 -*-
"""
【26秋】7年级讲义 · 长难句解析与必背表达整理（第 1–3 讲）
================================================================
依据同目录《【26秋】7年级讲义（第1-3讲）.docx》中的三篇语篇（第1讲 友谊阅读、
第2讲 亲情完形、第3讲 自我·家庭范文/语法填空），按「句 → 解析 → 必背表达 → 仿写」
四段式整理，版式参考样卡（◆句 / 🔑长难句解析 / 必背表达加粗 / 仿写挖空），
难度对齐 7 年级（七下高频句型与短语）。复用 jiangyi_lib 的青绿+暖橙现代版式。

运行：
    cd 26秋7年级讲义
    python3 build_changnanju.py
依赖：python-docx
"""
from jiangyi_lib import (
    Jiangyi, EA_HEAD, EA_BODY, EN_BODY, EN_SERIF,
    TEAL, TEAL_MID, ORANGE, RED, INK, GREY, WHITE,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---- 行内样式简写 ----
K   = {'bold': True, 'color': ORANGE, 'latin': EN_SERIF, 'ea': EA_HEAD}      # 关键短语（橙·加粗）
M   = {'bold': True, 'color': TEAL_MID, 'latin': EN_BODY, 'ea': EA_HEAD}     # 圈码 ①②③
EN  = {'latin': EN_SERIF, 'ea': EA_BODY}                                     # 英文常规
CNK = {'bold': True, 'color': ORANGE, 'ea': EA_HEAD}                         # 译文里的关键词（橙）
CN  = {'ea': EA_BODY}


# ============================================================
# 一张「长难句卡片」 = ◆句 + 英汉框 + 🔑解析 + 必背表达 + 仿写
# ============================================================
def sentence_card(J, n, en_segs, cn_segs, analysis, must, fanxie):
    # ◆ 句 N
    p = J._p(space_before=9, space_after=2, keep=True)
    r0 = p.add_run("◆ ")
    J._set = None
    from jiangyi_lib import _set_run_font
    _set_run_font(r0, size=13, bold=True, color=ORANGE)
    r1 = p.add_run(f"句 {n}")
    _set_run_font(r1, size=12.5, bold=True, color=TEAL, ea=EA_HEAD)

    # 英汉对照框（暖橙浅底，仿样卡的浅黄框）
    J.box([en_segs, cn_segs], kind="orange")

    # 🔑 长难句解析
    J.section_header("长难句解析", icon="🔑")
    for mark, title, subs in analysis:
        pp = J._p(space_before=2, space_after=1, indent=0.2, keep=True)
        rm = pp.add_run(f"{mark} ")
        _set_run_font(rm, size=11, bold=True, color=ORANGE)
        rt = pp.add_run(title)
        _set_run_font(rt, size=11, bold=True, color=INK, ea=EA_HEAD)
        for sub in subs:
            J.body(sub, space_after=1, indent=0.7, size=10)

    # 必背表达（标题加粗，下列条目一一对应、关键短语加粗）
    J.section_header("必背表达", icon="📘")
    for i, (exp, mean) in enumerate(must, 1):
        J.body([
            (f"{i}) ", {'bold': True, 'color': TEAL_MID, 'ea': EA_HEAD}),
            (exp, {'bold': True, 'color': INK, 'latin': EN_BODY, 'ea': EA_HEAD}),
            (f"　{mean}", {'ea': EA_BODY, 'size': 10}),
        ], space_after=1, indent=0.4, size=10.5)

    # ✏️ 仿写练习（挖空：给中文 + 括号提示核心句型/短语）
    J.section_header("仿写练习", icon="✏️")
    cn_prompt, hint, answer = fanxie
    J.body([
        ("翻译：", {'bold': True, 'color': RED, 'ea': EA_HEAD}),
        (cn_prompt, {'ea': EA_BODY}),
        (f"（提示：{hint}）", {'color': GREY, 'size': 9.5, 'ea': EA_BODY}),
    ], space_after=2, indent=0.2, size=10.5)
    # 作答横线
    pl = J._p(space_after=3, indent=0.2)
    rl = pl.add_run("✎ " + "_" * 52)
    _set_run_font(rl, size=10.5, color=GREY)
    return answer  # 收集到本讲参考答案


from jiangyi_lib import _set_run_font  # noqa: E402  (供卡片函数使用)


def lesson_intro(J, lesson_no, cn_title, source):
    """每讲小标题横幅 + 语篇出处说明。"""
    J.part_header(f"Lesson {lesson_no}", cn_title)
    J.body([
        ("语篇出处：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD, 'size': 10}),
        (source, {'color': GREY, 'size': 10, 'ea': EA_BODY}),
    ], space_after=4, size=10)


def answer_box(J, answers):
    J.answer_key("仿写练习 · 参考答案", [
        [(f"句{i} ", {'bold': True, 'color': ORANGE, 'ea': EA_HEAD, 'size': 9.5}),
         (a, {'latin': EN_BODY, 'size': 9.5, 'ea': EA_BODY})]
        for i, a in enumerate(answers, 1)
    ])


# ============================================================
# 内容数据
# ============================================================

def build():
    J = Jiangyi(title="【26秋】7年级讲义 · 长难句解析")

    # -------- 文档封面（首页）--------
    sec = J.doc.sections[0]
    J._set_footer(sec, "【26秋】7年级 · 长难句解析")
    sec.different_first_page_header_footer = False
    pt = J._p(space_before=40, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    rt = pt.add_run("长难句解析 与 必背表达整理")
    _set_run_font(rt, size=24, bold=True, color=TEAL, ea=EA_HEAD)
    ps = J._p(space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    rs = ps.add_run("【26秋】7 年级讲义 · 第 1–3 讲")
    _set_run_font(rs, size=14, bold=True, color=ORANGE, ea=EA_HEAD)
    pe = J._p(space_after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    re_ = pe.add_run("Long Sentences · Key Expressions · Imitation Writing")
    _set_run_font(re_, size=11, color=GREY, latin=EN_BODY)

    J.box([
        [("使用说明", {'bold': True, 'color': TEAL, 'ea': EA_HEAD})],
        [("①②③", {'bold': True, 'color': TEAL_MID, 'ea': EA_HEAD}),
         (" 为句中要点序号，与下方", {'ea': EA_BODY}),
         ("🔑 长难句解析", {'bold': True, 'color': INK, 'ea': EA_HEAD}),
         (" 一一对应；", {'ea': EA_BODY})],
        [("• ", {'color': ORANGE}),
         ("📘 必背表达", {'bold': True, 'color': INK, 'ea': EA_HEAD}),
         (" 标题与短语均加粗，逐条编号对应该句核心句型 / 短语；", {'ea': EA_BODY})],
        [("• ", {'color': ORANGE}),
         ("✏️ 仿写练习", {'bold': True, 'color': INK, 'ea': EA_HEAD}),
         (" 仿照上方讲解的句型 / 短语挖空，给中文并在括号内提示考点，文末附参考答案。", {'ea': EA_BODY})],
    ], kind="teal")

    # ========== 第 1 讲 ==========
    lesson_intro(J, 1, "友谊·品格（阅读：Friends are God's way of looking after us）",
                 "第1讲 Part 阅读精练《Friends are God's way of looking after us》【江苏期中】")
    ans1 = []
    ans1.append(sentence_card(
        J, 1,
        [("Sally says ", EN), ("①", M), ("finding friendship", K), (" is just ", EN),
         ("②", M), ("like planting", K), (" a tree.", EN)],
        [("莎莉说，", CN), ("①寻找友谊", CNK), ("就", CN), ("②像种", CNK), ("一棵树一样。", CN)],
        [
            ("①", "finding friendship —— 动名词短语作主语",
             [[("动名词（v.-ing）可作主语，表示“做……这件事”，谓语用单数：", CN), ("is", {'latin': EN_BODY, 'bold': True, 'color': TEAL})]]),
            ("②", "be (just) like doing sth. —— （正）像做某事一样",
             [[("like 是介词，后接名词或动名词 ", CN), ("doing", {'latin': EN_BODY, 'bold': True, 'color': TEAL}), ("；just 加强“正像”。", CN)]]),
        ],
        [("say (that) ...", "说……"),
         ("be like doing sth.", "像做某事一样"),
         ("plant a tree", "种树")],
        ("学英语就像交朋友一样。", "be like doing", "Learning English is just like making friends."),
    ))
    ans1.append(sentence_card(
        J, 2,
        [("A good friend should ", EN), ("①", M), ("listen to", K), (" your complaints and ", EN),
         ("②", M), ("do his or her best", K), (" to help you.", EN)],
        [("一个好朋友应该", CN), ("①倾听", CNK), ("你的抱怨，并", CN), ("②尽他/她最大的努力", CNK), ("来帮助你。", CN)],
        [
            ("①", "listen to sth. —— 倾听……",
             [[("listen 是不及物动词，“听某物”须加介词 ", CN), ("to", {'latin': EN_BODY, 'bold': True, 'color': TEAL}), ("。", CN)]]),
            ("②", "do one's best to do sth. —— 尽某人最大努力去做某事",
             [[("one's 随主语变化（my/your/his/her…）；后接 ", CN), ("to do", {'latin': EN_BODY, 'bold': True, 'color': TEAL}), (" 不定式。", CN)]]),
        ],
        [("listen to", "倾听"),
         ("do one's best to do sth.", "尽力做某事"),
         ("help sb. (to) do sth.", "帮助某人做某事")],
        ("我会尽我最大的努力来帮助你。", "do one's best", "I will do my best to help you."),
    ))
    ans1.append(sentence_card(
        J, 3,
        [("①", M), ("When", K), (" there are no other people around you, ", EN),
         ("②", M), ("have an honest talk", K), (".", EN)],
        [("①当", CNK), ("你周围没有其他人的时候，", CN), ("②进行一次坦诚的谈话", CNK), ("。", CN)],
        [
            ("①", "when 引导【时间状语从句】",
             [[("主句：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("have an honest talk（祈使句）", EN)],
              [("引导词：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("when（当……时）", EN)],
              [("从句：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("there are no other people around you（there be 句型）", EN)]]),
            ("②", "have a/an + 形容词 + talk —— 进行一次……的谈话",
             [[("have 表“进行/从事”，如 have a rest / have a talk。", CN)]]),
        ],
        [("when", "当……时候"),
         ("there be", "（某处）有……"),
         ("have a talk", "谈一谈；around 在……周围")],
        ("当我遇到麻烦时，我会和老师谈一谈。", "when … have a talk", "When I am in trouble, I have a talk with my teacher."),
    ))
    ans1.append(sentence_card(
        J, 4,
        [("Remember ", EN), ("①", M), ("that", K), (" friendship is ", EN),
         ("②", M), ("the most important", K), (" thing in your life.", EN)],
        [("记住，", CN), ("①友谊是", CN), ("你生命中", CN), ("②最重要的", CNK), ("东西。", CN)],
        [
            ("①", "that 引导【宾语从句】",
             [[("主句：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("Remember …（祈使句）", EN)],
              [("that 从句作 remember 的宾语，口语中 ", CN), ("that", {'latin': EN_BODY, 'bold': True, 'color': TEAL}), (" 可省略。", CN)]]),
            ("②", "the most important —— 形容词最高级",
             [[("important 为多音节词，最高级加 most：", CN), ("the most + 多音节形容词", {'latin': EN_BODY, 'bold': True, 'color': TEAL})]]),
        ],
        [("remember that ...", "记住……"),
         ("the most important", "最重要的"),
         ("in one's life", "在某人一生中")],
        ("请记住，健康是最重要的东西。", "the most important", "Please remember that health is the most important thing."),
    ))
    answer_box(J, ans1)

    # ========== 第 2 讲 ==========
    lesson_intro(J, 2, "亲情·陪伴（完形：Remember to call your parents）",
                 "第2讲 Part 完形精练《Remember to call your parents》【广东期中】")
    ans2 = []
    ans2.append(sentence_card(
        J, 1,
        [("My parents like the Mid-Autumn Festival very much ", EN), ("①", M), ("because", K),
         (" my aunt and uncle ", EN), ("②", M), ("come back from", K), (" Beijing to see them.", EN)],
        [("我父母非常喜欢中秋节，", CN), ("①因为", CNK), ("我的姑姑和叔叔", CN), ("②从北京回来", CNK), ("看他们。", CN)],
        [
            ("①", "because 引导【原因状语从句】",
             [[("主句：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("my parents like the festival very much", EN)],
              [("引导词：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("because（因为，回答 Why）", EN)]]),
            ("②", "come back from + 地点 + to do sth. —— 从某地回来做某事",
             [[("to see them 是不定式作", CN), ("目的状语", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("（为了……）。", CN)]]),
        ],
        [("because", "因为"),
         ("come back from", "从……回来"),
         ("… very much", "非常……")],
        ("我喜欢周末，因为我可以和家人待在一起。", "because", "I like weekends because I can stay with my family."),
    ))
    ans2.append(sentence_card(
        J, 2,
        [("My parents are old, ", EN), ("①", M), ("so", K), (" they ", EN),
         ("②", M), ("can't get close to", K), (" the telephone quickly.", EN)],
        [("我的父母年纪大了，", CN), ("①所以", CNK), ("他们", CN), ("②不能很快地靠近", CNK), ("电话。", CN)],
        [
            ("①", "so —— 表“所以”的并列连词（连接结果）",
             [[("避坑：英语里 ", {'color': RED}), ("because 与 so 不能同时使用", {'bold': True, 'color': RED, 'ea': EA_HEAD}), ("，二者用其一。", {'color': RED})]]),
            ("②", "get close to sth. —— 靠近某物",
             [[("close 此处是形容词“近的”；get + 形容词表“变得……”。", CN)]]),
        ],
        [("so", "所以"),
         ("get close to", "靠近"),
         ("can't do sth. quickly", "不能很快做某事")],
        ("天气很冷，所以我们待在家里。", "so", "It is cold, so we stay at home."),
    ))
    ans2.append(sentence_card(
        J, 3,
        [("I only ", EN), ("①", M), ("want to give", K), (" them ", EN),
         ("②", M), ("enough time to answer", K), (" the telephone.", EN)],
        [("我只是", CN), ("①想给", CNK), ("他们", CN), ("②足够的时间来接", CNK), ("电话。", CN)],
        [
            ("①", "want to do sth. —— 想要做某事",
             [[("want 后接不定式 ", CN), ("to do", {'latin': EN_BODY, 'bold': True, 'color': TEAL}), ("，不接 doing。", CN)]]),
            ("②", "give sb. + enough time + to do sth. —— 给某人足够的时间做某事",
             [[("enough 修饰名词放其", CN), ("前", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("：enough time；后接 to do 表用途。", CN)]]),
        ],
        [("want to do sth.", "想要做某事"),
         ("give sb. time to do sth.", "给某人时间做某事"),
         ("answer the telephone", "接电话")],
        ("妈妈给了我足够的时间来完成作业。", "enough time to do", "My mother gave me enough time to finish my homework."),
    ))
    ans2.append(sentence_card(
        J, 4,
        [("Please always ", EN), ("①", M), ("remember to call", K), (" your parents ", EN),
         ("②", M), ("at any time", K), (".", EN)],
        [("请永远", CN), ("①记得给", CNK), ("你的父母打电话，", CN), ("②在任何时候", CNK), ("。", CN)],
        [
            ("①", "remember to do sth. —— 记得（去）做某事（事还没做）",
             [[("对比：", CN), ("remember doing sth.", {'latin': EN_BODY, 'bold': True, 'color': TEAL}), (" 记得做过某事（事已做）。", CN)]]),
            ("②", "at any time —— 在任何时候",
             [[("any 用于肯定句意为“任何”。", CN)]]),
        ],
        [("remember to do sth.", "记得去做某事"),
         ("call sb.", "给某人打电话"),
         ("at any time", "在任何时候")],
        ("请记得每天给奶奶打电话。", "remember to do", "Please remember to call your grandma every day."),
    ))
    answer_box(J, ans2)

    # ========== 第 3 讲 ==========
    lesson_intro(J, 3, "自我·家庭（范文 My Family / 自我介绍 + 语法填空）",
                 "第3讲 Part 满分作文《My Family》《自我介绍》与语法填空【江苏月考】")
    ans3 = []
    ans3.append(sentence_card(
        J, 1,
        [("I like ", EN), ("①", M), ("playing basketball", K), (" very much, ", EN),
         ("②", M), ("because", K), (" it can ", EN), ("③", M), ("help me keep healthy", K), (".", EN)],
        [("我非常喜欢", CN), ("①打篮球", CNK), ("，", CN), ("②因为", CNK), ("它能", CN), ("③帮助我保持健康", CNK), ("。", CN)],
        [
            ("①", "like doing sth. —— 喜欢做某事（表习惯爱好）", []),
            ("②", "because 引导【原因状语从句】", []),
            ("③", "help sb. (to) do sth. —— 帮助某人做某事",
             [[("to 常省略；keep + 形容词：", CN), ("keep healthy", {'latin': EN_BODY, 'bold': True, 'color': TEAL}), (" 保持健康。", CN)]]),
        ],
        [("like doing sth.", "喜欢做某事"),
         ("help sb. do sth.", "帮助某人做某事"),
         ("keep healthy", "保持健康")],
        ("我喜欢跑步，因为它能帮助我保持健康。", "help sb. do", "I like running because it can help me keep healthy."),
    ))
    ans3.append(sentence_card(
        J, 2,
        [("My mother always tells me ", EN), ("①", M), ("not to give up", K), (" ", EN),
         ("②", M), ("when", K), (" I'm in trouble.", EN)],
        [("②当我遇到困难时", CNK), ("，我妈妈总是告诉我", CN), ("①不要放弃", CNK), ("。", CN)],
        [
            ("①", "tell sb. (not) to do sth. —— 告诉某人（不要）做某事",
             [[("否定式在 to 前加 not：", CN), ("not to do", {'latin': EN_BODY, 'bold': True, 'color': TEAL}), ("。", CN)]]),
            ("②", "when 引导【时间状语从句】",
             [[("主句：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("my mother tells me not to give up", EN)],
              [("从句：", {'bold': True, 'color': TEAL, 'ea': EA_HEAD}), ("when I'm in trouble（be in trouble 处于困境）", EN)]]),
        ],
        [("tell sb. to do sth.", "告诉某人做某事"),
         ("give up", "放弃"),
         ("be in trouble", "处于困境")],
        ("当我遇到困难时，老师总是告诉我不要害怕。", "tell sb. not to do", "When I am in trouble, my teacher always tells me not to be afraid."),
    ))
    ans3.append(sentence_card(
        J, 3,
        [("She hopes ", EN), ("①", M), ("to be", K), (" a singer ", EN),
         ("②", M), ("when she grows up", K), (".", EN)],
        [("她希望", CN), ("②长大后", CNK), ("①成为", CNK), ("一名歌手。", CN)],
        [
            ("①", "hope to do sth. —— 希望做某事",
             [[("hope 后接不定式 to do，不接 doing。", CN)]]),
            ("②", "when she grows up —— 时间状语从句",
             [[("grow up 长大；从句用一般现在时表将来。", CN)]]),
        ],
        [("hope to do sth.", "希望做某事"),
         ("grow up", "长大"),
         ("be good at (doing) sth.", "擅长（做）某事")],
        ("我希望长大后成为一名老师。", "hope to do, grow up", "I hope to be a teacher when I grow up."),
    ))
    answer_box(J, ans3)

    out = "【26秋】7年级讲义·长难句解析与表达整理（第1-3讲）.docx"
    J.save(out)
    print("saved:", out)


if __name__ == "__main__":
    build()
