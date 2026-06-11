#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
【26春】7年级英语 L16 正式课课件生成器
4:3 / 77页 / 答案点击出现动画(单击①答案高亮✓ → 单击②考点解析条滑入)
版式代号对应设计稿:封面/目录/章节/知识/例题/解析/互动/小结/作业/结束
注:23秋母版源文件缺失,本稿用同结构自绘版式占位,待源文件到位后重绑母版。
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ---------- 设计系统(23秋直播课风格近似) ----------
NAVY    = RGBColor(0x1F, 0x38, 0x64)   # 深蓝主色
ORANGE  = RGBColor(0xFF, 0x7A, 0x00)   # 橙主色
ORANGE_L= RGBColor(0xFF, 0xF3, 0xE6)   # 浅橙底
BLUE    = RGBColor(0x2E, 0x75, 0xB6)   # 蓝辅色
BLUE_L  = RGBColor(0xEC, 0xF3, 0xFB)   # 浅蓝底
GREEN   = RGBColor(0x00, 0x9E, 0x4F)   # 答案绿
GREEN_L = RGBColor(0xE6, 0xF7, 0xEE)
RED     = RGBColor(0xE5, 0x39, 0x35)   # 警示红
YELLOW  = RGBColor(0xFF, 0xF2, 0x00)   # 标黄
YELLOW_L= RGBColor(0xFF, 0xFB, 0xD9)
GRAY    = RGBColor(0x59, 0x5959 >> 8, 0x59)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
DARK    = RGBColor(0x26, 0x26, 0x26)
FONT    = "微软雅黑"

PAGE_W, PAGE_H = Inches(10), Inches(7.5)

prs = Presentation()
prs.slide_width  = PAGE_W
prs.slide_height = PAGE_H
BLANK = prs.slide_layouts[6]

TOTAL = 77
_page = [0]

# ---------- 基础工具 ----------
def _set_font(run, size=18, bold=False, color=DARK, italic=False, name=FONT):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color; f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = etree.SubElement(rPr, qn(tag))
        e.set("typeface", name)

def rich(p, text, size=18, bold=False, color=DARK, bold_color=None):
    """支持 **加粗高亮** 标记;bold_color 默认橙色"""
    hl = bold_color or ORANGE
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        r = p.add_run(); r.text = seg
        if i % 2 == 1:
            _set_font(r, size=size, bold=True, color=hl)
        else:
            _set_font(r, size=size, bold=bold, color=color)

def box(slide, x, y, w, h, fill=WHITE, line=None, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, shadow=False, name=None):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = radius
        except Exception: pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if name: sp.name = name
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return sp

def txt(slide, x, y, w, h, name=None):
    sp = slide.shapes.add_textbox(x, y, w, h)
    if name: sp.name = name
    tf = sp.text_frame
    tf.word_wrap = True
    return sp

def para(tf, text="", size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT,
         bold_color=None, space_after=4, first=False, line_spacing=None):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if line_spacing: p.line_spacing = line_spacing
    if text:
        rich(p, text, size=size, bold=bold, color=color, bold_color=bold_color)
    return p

def chip(slide, x, y, w, h, text, fill=ORANGE, color=WHITE, size=14, bold=True):
    sp = box(slide, x, y, w, h, fill=fill, radius=0.5)
    para(sp.text_frame, text, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER, first=True, space_after=0)
    return sp

# ---------- 公共版式件 ----------
def footer(slide):
    _page[0] += 1
    n = _page[0]
    tf = txt(slide, Inches(0.3), Inches(7.08), Inches(4), Inches(0.35)).text_frame
    para(tf, "26春 · 七年级英语 · L16", size=10, color=GRAY, first=True, space_after=0)
    tf2 = txt(slide, Inches(8.5), Inches(7.08), Inches(1.2), Inches(0.35)).text_frame
    para(tf2, f"{n} / {TOTAL}", size=10, color=GRAY, align=PP_ALIGN.RIGHT, first=True, space_after=0)
    return n

def header(slide, tag, title, source=None, tag_fill=ORANGE):
    box(slide, 0, 0, PAGE_W, Inches(0.16), fill=tag_fill, shape=MSO_SHAPE.RECTANGLE)
    chip(slide, Inches(0.35), Inches(0.38), Inches(1.25), Inches(0.46), tag, fill=tag_fill, size=15)
    tf = txt(slide, Inches(1.78), Inches(0.30), Inches(6.2), Inches(0.62)).text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, title, size=23, bold=True, color=NAVY, first=True, space_after=0)
    if source:
        chip(slide, Inches(8.05), Inches(0.42), Inches(1.6), Inches(0.40), source,
             fill=BLUE_L, color=BLUE, size=12)
    ln = box(slide, Inches(0.35), Inches(1.02), Inches(9.3), Pt(2.2), fill=tag_fill, shape=MSO_SHAPE.RECTANGLE)
    return Inches(1.25)  # 内容起始 y

def new_slide():
    return prs.slides.add_slide(BLANK)

# ---------- 动画(p:timing 注入) ----------
# 形状命名约定:
#   auto:K|effect  → 随页自动浮现,K=级联步序(0,1,2...),每步延迟250ms
#   clickN|effect  → 第N次单击触发出现
# effect: fade(淡入) / wipe(自左擦入) / fly(自底飞入)
NSMAP_P = "http://schemas.openxmlformats.org/presentationml/2006/main"

_PRESETS = {"fade": (10, 0), "wipe": (22, 8), "fly": (2, 4)}

def _effect_par(ids, spid, effect, node_type, delay=0):
    pid, psub = _PRESETS[effect]
    e1, e2 = next(ids), next(ids)
    parts = [f"""<p:set>
    <p:cBhvr>
     <p:cTn id="{e2}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
     <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
     <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
    </p:cBhvr>
    <p:to><p:strVal val="visible"/></p:to>
   </p:set>"""]
    if effect in ("fade", "wipe"):
        filt = "fade" if effect == "fade" else "wipe(right)"
        e3 = next(ids)
        parts.append(f"""<p:animEffect transition="in" filter="{filt}">
    <p:cBhvr><p:cTn id="{e3}" dur="500"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr>
   </p:animEffect>""")
    else:  # fly: 自底飞入(ppt_x 不变,ppt_y 从屏外到位)
        a1, a2 = next(ids), next(ids)
        parts.append(f"""<p:anim calcmode="lin" valueType="num">
    <p:cBhvr additive="base">
     <p:cTn id="{a1}" dur="500" fill="hold"/>
     <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
     <p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst>
    </p:cBhvr>
    <p:tavLst>
     <p:tav tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav>
     <p:tav tm="100000"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav>
    </p:tavLst>
   </p:anim>
   <p:anim calcmode="lin" valueType="num">
    <p:cBhvr additive="base">
     <p:cTn id="{a2}" dur="500" fill="hold"/>
     <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
     <p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst>
    </p:cBhvr>
    <p:tavLst>
     <p:tav tm="0"><p:val><p:strVal val="1+#ppt_h/2"/></p:val></p:tav>
     <p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav>
    </p:tavLst>
   </p:anim>""")
    return f"""<p:par xmlns:p="{NSMAP_P}">
 <p:cTn id="{e1}" presetID="{pid}" presetClass="entr" presetSubtype="{psub}" fill="hold" grpId="0" nodeType="{node_type}">
  <p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>
  <p:childTnLst>{''.join(parts)}</p:childTnLst>
 </p:cTn>
</p:par>"""

def _wrap_group(ids, inner, trigger_delay):
    g1, g2 = next(ids), next(ids)
    return f"""<p:par xmlns:p="{NSMAP_P}">
 <p:cTn id="{g1}" fill="hold">
  <p:stCondLst><p:cond delay="{trigger_delay}"/></p:stCondLst>
  <p:childTnLst>
   <p:par>
    <p:cTn id="{g2}" fill="hold">
     <p:stCondLst><p:cond delay="0"/></p:stCondLst>
     <p:childTnLst>{''.join(inner)}</p:childTnLst>
    </p:cTn>
   </p:par>
  </p:childTnLst>
 </p:cTn>
</p:par>"""

def apply_click_animations(slide):
    auto, clicks, bld = {}, {}, []
    for sp in slide.shapes:
        nm = sp.name or ""
        if nm.startswith("auto:"):
            head = nm.split("|")
            step = int(head[0][5:])
            eff = head[1] if len(head) > 1 else "fade"
            is_frame = sp._element.tag.endswith("graphicFrame")
            auto.setdefault(step, []).append((sp.shape_id, eff, is_frame))
        elif nm.startswith("click"):
            head = nm.split("|")
            idx = int(head[0][5:])
            eff = head[1] if len(head) > 1 else "fade"
            is_frame = sp._element.tag.endswith("graphicFrame")
            clicks.setdefault(idx, []).append((sp.shape_id, eff, is_frame))
    if not auto and not clicks:
        return
    counter = iter(range(3, 100000))
    seq_children = []
    if auto:
        inner, first = [], True
        for step in sorted(auto):
            for spid, eff, isf in auto[step]:
                nt = "afterEffect" if first else "withEffect"
                inner.append(_effect_par(counter, spid, eff, nt, delay=step * 250))
                bld.append((spid, isf))
                first = False
        seq_children.append(_wrap_group(counter, inner, "0"))
    for idx in sorted(clicks):
        inner = []
        for j, (spid, eff, isf) in enumerate(clicks[idx]):
            nt = "clickEffect" if j == 0 else "withEffect"
            inner.append(_effect_par(counter, spid, eff, nt))
            bld.append((spid, isf))
        seq_children.append(_wrap_group(counter, inner, "indefinite"))
    bld_xml = "".join(
        (f'<p:bldGraphic spid="{spid}" grpId="0"><p:bldAsOne/></p:bldGraphic>'
         if isf else f'<p:bldP spid="{spid}" grpId="0"/>')
        for spid, isf in bld)
    timing = f"""<p:timing xmlns:p="{NSMAP_P}">
 <p:tnLst>
  <p:par>
   <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
    <p:childTnLst>
     <p:seq concurrent="1" nextAc="seek">
      <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
       <p:childTnLst>{''.join(seq_children)}</p:childTnLst>
      </p:cTn>
      <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
      <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
     </p:seq>
    </p:childTnLst>
   </p:cTn>
  </p:par>
 </p:tnLst>
 <p:bldLst>{bld_xml}</p:bldLst>
</p:timing>"""
    slide._element.append(etree.fromstring(timing))

# ---------- 版式:封面/章节/目录/结束 ----------
def slide_cover(title_lines, subtitle, tag_text):
    s = new_slide()
    box(s, 0, 0, PAGE_W, PAGE_H, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    box(s, 0, Inches(5.9), PAGE_W, Inches(0.12), fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    box(s, Inches(8.6), Inches(0.0), Inches(1.4), Inches(7.5), fill=RGBColor(0x27,0x44,0x77), shape=MSO_SHAPE.RECTANGLE)
    c = chip(s, Inches(0.7), Inches(0.9), Inches(2.6), Inches(0.5), tag_text, fill=ORANGE, size=16)
    c.name = "auto:0|fade"
    tf = txt(s, Inches(0.7), Inches(2.2), Inches(8.6), Inches(2.6), name="auto:1|fly").text_frame
    for i, line in enumerate(title_lines):
        para(tf, line, size=40, bold=True, color=WHITE, first=(i == 0), space_after=10)
    tf2 = txt(s, Inches(0.7), Inches(4.9), Inches(8.6), Inches(0.7), name="auto:2|fade").text_frame
    para(tf2, subtitle, size=20, color=RGBColor(0xCF,0xDD,0xF2), first=True)
    tf3 = txt(s, Inches(0.7), Inches(6.3), Inches(8), Inches(0.5), name="auto:3|fade").text_frame
    para(tf3, "主讲老师:________        正式课", size=16, color=RGBColor(0x9F,0xB5,0xD5), first=True)
    footer(s); apply_click_animations(s); return s

def slide_section(part_no, title_en, title_cn, items=None, fill=NAVY):
    s = new_slide()
    box(s, 0, 0, PAGE_W, PAGE_H, fill=fill, shape=MSO_SHAPE.RECTANGLE)
    box(s, 0, Inches(2.0), Inches(0.25), Inches(3.5), fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    tf = txt(s, Inches(0.8), Inches(1.7), Inches(8.4), Inches(1.0), name="auto:0|fly").text_frame
    para(tf, part_no, size=44, bold=True, color=ORANGE, first=True)
    tf2 = txt(s, Inches(0.8), Inches(2.9), Inches(8.6), Inches(1.0), name="auto:1|fade").text_frame
    para(tf2, title_en, size=30, bold=True, color=WHITE, first=True, space_after=6)
    para(tf2, title_cn, size=24, bold=True, color=RGBColor(0xCF,0xDD,0xF2))
    if items:
        tf3 = txt(s, Inches(0.85), Inches(4.6), Inches(8.4), Inches(1.8), name="auto:2|fade").text_frame
        for i, it in enumerate(items):
            para(tf3, "·  " + it, size=17, color=RGBColor(0xBF,0xD0,0xE8), first=(i == 0), space_after=6)
    footer(s); apply_click_animations(s); return s

def slide_end():
    s = new_slide()
    box(s, 0, 0, PAGE_W, PAGE_H, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    tf = txt(s, Inches(0.5), Inches(2.3), Inches(9), Inches(1.2), name="auto:0|fly").text_frame
    para(tf, "谢谢观看  Thanks!", size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
    b = box(s, Inches(1.5), Inches(4.2), Inches(7), Inches(1.5), fill=RGBColor(0x27,0x44,0x77), name="auto:1|fade")
    tf2 = b.text_frame
    para(tf2, "下讲预告", size=16, bold=True, color=ORANGE, align=PP_ALIGN.CENTER, first=True, space_after=8)
    para(tf2, "L17 期末复习①(篇章:语法填空+阅读+完形)", size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    footer(s); apply_click_animations(s); return s

# ---------- 版式:知识/解析/小结(标题+内容块流式) ----------
def content_blocks(s, y0, blocks, auto=True, step0=0):
    """blocks: ('p',text,size) / ('box',title,lines,fill,line) / ('table',rows,widths,heights) / ('gap',h)"""
    y = y0
    step = step0
    def anm():
        nonlocal step
        n = f"auto:{step}|fade" if auto else None
        step += 1
        return n
    for blk in blocks:
        kind = blk[0]
        if kind == "gap":
            y += Inches(blk[1])
        elif kind == "p":
            text, size = blk[1], (blk[2] if len(blk) > 2 else 17)
            kw = blk[3] if len(blk) > 3 else {}
            n_lines = max(1, len(text) * (size / 2 + 2) // 620 + 1) if False else text.count("\n") + 1
            h = Inches(0.34) * n_lines * (size / 17.0)
            sp = txt(s, Inches(0.5), y, Inches(9.0), h, name=anm())
            tf = sp.text_frame
            for i, ln in enumerate(text.split("\n")):
                para(tf, ln, size=size, first=(i == 0), space_after=4, **kw)
            y += h + Inches(0.06)
        elif kind == "box":
            title, lines, fill, line_c = blk[1], blk[2], blk[3], blk[4]
            h = blk[5] if len(blk) > 5 else Inches(0.42 + 0.34 * len(lines) + (0.34 if title else 0))
            b = box(s, Inches(0.5), y, Inches(9.0), h, fill=fill, line=line_c, name=anm())
            tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            first = True
            if title:
                para(tf, title, size=17, bold=True, color=NAVY, first=True, space_after=6)
                first = False
            for ln in lines:
                para(tf, ln, size=16, first=first, space_after=4)
                first = False
            y += h + Inches(0.14)
        elif kind == "table":
            rows, widths = blk[1], blk[2]
            row_h = blk[3] if len(blk) > 3 else 0.42
            total_w = sum(widths)
            x0 = (10 - total_w) / 2
            gtbl = s.shapes.add_table(len(rows), len(widths), Inches(x0), y, Inches(total_w), Inches(row_h * len(rows)))
            nm = anm()
            if nm: gtbl.name = nm
            tbl = gtbl.table
            for ci, wd in enumerate(widths):
                tbl.columns[ci].width = Inches(wd)
            for ri, row in enumerate(rows):
                tbl.rows[ri].height = Inches(row_h)
                for ci, cell_text in enumerate(row):
                    c = tbl.cell(ri, ci)
                    c.vertical_anchor = MSO_ANCHOR.MIDDLE
                    c.margin_left = c.margin_right = Inches(0.06)
                    c.margin_top = c.margin_bottom = Inches(0.02)
                    tf = c.text_frame; tf.word_wrap = True
                    hdr = (ri == 0)
                    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                    rich(p, cell_text, size=14 if not hdr else 15,
                         bold=hdr, color=WHITE if hdr else DARK,
                         bold_color=WHITE if hdr else ORANGE)
                    c.fill.solid()
                    c.fill.fore_color.rgb = BLUE if hdr else (WHITE if ri % 2 else BLUE_L)
                    if hdr: c.fill.fore_color.rgb = NAVY
            y += Inches(row_h * len(rows)) + Inches(0.15)
    return y

def slide_know(tag, title, blocks, source=None, tag_fill=BLUE):
    s = new_slide()
    y0 = header(s, tag, title, source=source, tag_fill=tag_fill)
    content_blocks(s, y0 + Inches(0.1), blocks)
    footer(s); apply_click_animations(s); return s

# ---------- 版式:例题(题干+选项+点击答案+点击解析条) ----------
def slide_ex(tag, title, source, stem, options=None, answer=None, ans_label=None,
             analysis=None, stem_h=None, opts_cols=2, extra_clicks=None, opt_y=None):
    """
    options: ["A. ...","B. ...","C. ...","D. ..."]  answer: 正确项序号0-3
    ans_label: 文本答案(填空/转换题) — 单击①出现
    analysis: 解析条文本 — 单击②滑入
    extra_clicks: [(click_idx,text), ...] 分步出现的答案条(P33等)
    """
    s = new_slide()
    y0 = header(s, tag, title, source=source, tag_fill=ORANGE)
    sh = stem_h or Inches(1.15)
    sb = box(s, Inches(0.5), y0 + Inches(0.12), Inches(9.0), sh, fill=WHITE, line=RGBColor(0xD9,0xD9,0xD9),
             name="auto:0|fade")
    tf = sb.text_frame
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, ln in enumerate(stem.split("\n")):
        para(tf, ln, size=19, color=DARK, first=(i == 0), space_after=6, bold_color=BLUE)
    y = y0 + Inches(0.12) + sh + Inches(0.2)
    if opt_y: y = opt_y
    if options:
        long_opts = max(len(o) for o in options) > 26 or opts_cols == 1
        if long_opts:
            ow, oh = Inches(8.6), Inches(0.52)
            positions = [(Inches(0.6), y + (oh + Inches(0.12)) * i) for i in range(len(options))]
        else:
            ow, oh = Inches(4.25), Inches(0.55)
            positions = [(Inches(0.6) + (ow + Inches(0.3)) * (i % 2),
                          y + (oh + Inches(0.16)) * (i // 2)) for i in range(len(options))]
        for i, (ox, oy) in enumerate(positions):
            ob = box(s, ox, oy, ow, oh, fill=BLUE_L, line=None, name=f"auto:{i+1}|fade")
            para(ob.text_frame, options[i], size=17, color=DARK, first=True, space_after=0)
            if i == answer:
                # 单击①:正确项绿框高亮 + ✓
                hb = box(s, ox - Inches(0.04), oy - Inches(0.04), ow + Inches(0.08), oh + Inches(0.08),
                         fill=None, line=GREEN, line_w=2.5, name="click1|fade")
                ck = box(s, ox + ow - Inches(0.5), oy + Inches(0.05), Inches(0.45), Inches(0.45),
                         fill=GREEN, shape=MSO_SHAPE.OVAL, name="click1|fade")
                para(ck.text_frame, "✓", size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True, space_after=0)
        y = positions[-1][1] + oh + Inches(0.2)
        if answer is not None:
            letter = "ABCD"[answer]
            ansc = chip(s, Inches(7.6), y, Inches(1.9), Inches(0.5), f"答案 · {letter}", fill=GREEN, size=16)
            ansc.name = "click1|fade"
    if ans_label:
        ab = box(s, Inches(0.6), y, Inches(8.8), Inches(0.62), fill=GREEN_L, line=GREEN, line_w=2, name="click1|fade")
        tfa = ab.text_frame
        p = tfa.paragraphs[0]
        r = p.add_run(); r.text = "答案  "; _set_font(r, size=15, bold=True, color=GREEN)
        rich(p, ans_label, size=18, bold=True, color=GREEN, bold_color=ORANGE)
        y += Inches(0.78)
    if extra_clicks:
        for idx, textv in extra_clicks:
            ab = box(s, Inches(0.6), y, Inches(8.8), Inches(0.6), fill=GREEN_L, line=GREEN, line_w=1.5,
                     name=f"click{idx}|fade")
            para(ab.text_frame, textv, size=17, bold=True, color=GREEN, first=True, space_after=0, bold_color=ORANGE)
            y += Inches(0.74)
    if analysis:
        ay = Inches(6.0)
        bar = box(s, Inches(0.5), ay, Inches(9.0), Inches(0.95), fill=ORANGE_L, line=ORANGE, line_w=1.5,
                  name="click2|wipe")
        tfb = bar.text_frame
        p = tfb.paragraphs[0]
        r = p.add_run(); r.text = "考点  "; _set_font(r, size=15, bold=True, color=ORANGE)
        rich(p, analysis.split("\n")[0], size=15, color=DARK)
        for ln in analysis.split("\n")[1:]:
            para(tfb, ln, size=15, color=DARK, space_after=2)
    footer(s); apply_click_animations(s); return s
